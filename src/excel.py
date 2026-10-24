import os
from openpyxl import Workbook, load_workbook
import logging 
import yaml

log_dir = "logs"
os.makedirs(log_dir, exist_ok=True)

logger=logging.getLogger("excel")
logger.setLevel("DEBUG")

console_handler=logging.StreamHandler()
console_handler.setLevel("DEBUG")

log_file_path=os.path.join(log_dir,"excel_file.log")
file_handler=logging.FileHandler(log_file_path)
file_handler.setLevel("DEBUG")

formatter = logging.Formatter('%(asctime)s - %(name)s - %(levelname)s - %(message)s')
file_handler.setFormatter(formatter)
console_handler.setFormatter(formatter)

logger.addHandler(file_handler)
logger.addHandler(console_handler)



def save_to_excel(data_dict, user_fields, file_path="output.xlsx"):
    try:

        if os.path.exists(file_path):
            # loading existing workbook
            wb = load_workbook(file_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(user_fields)  # header

        row = ws.max_row + 1
        for col, field in enumerate(user_fields, start=1):
            ws.cell(row=row, column=col, value=data_dict[field])

        wb.save(file_path)
        logger.debug("files saved successfully")
        
    except Exception as e:
        logger.error("Saving to excel failed %s",e)

